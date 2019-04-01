from django.shortcuts import render, redirect, reverse
from django.core import serializers
from django.core.serializers.json import DjangoJSONEncoder
from django.views import View
from django.http import JsonResponse, HttpResponseRedirect, HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils.datastructures import MultiValueDictKeyError
from django.utils.decorators import method_decorator
from guardian.models import UserObjectPermission
from guardian.decorators import *
from .forms import *
from .models import *
from users.models import User
from dispatch.models import Course
import subprocess
import random
import time
import json
import xlrd
import xlwt
import docx
from docxtpl import DocxTemplate
import re
try:
	import pythoncom
except ImportError as e:
	print(e)
try:
    from comtypes import client
except ImportError:
    client=None
import os
import numpy.matlib
import numpy as np


def is_teacher(func):
    """
    教师权限认证函数装饰器
    """
    def auth(request, *args, **kwargs):
        is_teacher = User.objects.get(user_id=request.user.user_id).is_teacher
        if not is_teacher:
            return redirect(reverse("users_login"))
        return func(request, *args, **kwargs)
    return auth


class Auth(View):
    """
    教师权限认证类装饰器
    """

    def dispatch(self, request, *args, **kwargs):
        is_teacher = User.objects.get(user_id=request.user.user_id).is_teacher
        if not is_teacher:
            return redirect(reverse("users_login"))
        return super(Auth, self).dispatch(request, *args, **kwargs)


class LoginRequiredMixin(object):
    """
    登录认证类装饰器
    """

    @method_decorator(login_required(login_url="/accounts/login/"))
    def dispatch(self, request, *args, **kwargs):
        return super(LoginRequiredMixin, self).dispatch(request, *args, **kwargs)


@is_teacher
def exam_auto_generate_page(request):
    """
    返回随机生成用页面
    """
    if request.method == 'POST':

        def case1():
            raw = Choice.objects.filter(
                courseID=courseID, is_single=True).values_list('sec', flat=True)
            clean_secs = sorted(list(set(raw)))
            return clean_secs

        def case2():
            raw = Choice.objects.filter(
                courseID=courseID, is_single=False).values_list('sec', flat=True)
            clean_secs = sorted(list(set(raw)))
            return clean_secs

        def case3():
            raw = Judge.objects.filter(
                courseID=courseID).values_list('sec', flat=True)
            clean_secs = sorted(list(set(raw)))
            return clean_secs

        def case4():

            raw = S_answer.objects.filter(
                courseID=courseID).values_list('sec', flat=True)
            clean_secs = sorted(list(set(raw)))
            return clean_secs

        def case5():
            raw = Blank.objects.filter(
                courseID=courseID).values_list('sec', flat=True)
            clean_secs = sorted(list(set(raw)))
            return clean_secs

        switch = {
            's_choice': case1,
            'm_choice': case2,
            'judge': case3,
            's_answer': case4,
            'blank': case5,
        }

        courseID = request.POST['courseID']
        e_type = ['s_choice', 'm_choice', 'judge', 's_answer', 'blank']
        secs = {}
        for e in e_type:
            secs[e] = switch[e]()
        print(secs)
        return render(request, 'partials/auto_gen.html', {'secs': secs})


@is_teacher
def exam_auto_generate(request):
    if request.method == 'POST':

        def pickle_in_random(secs_length, secs, count, e_type, id_list, per):
            """
            secs: 所有章节
            secs_length:各章节试题数
            count:需要随机组合的试题数
            e_type:当前题型
            id_list: 当前题型的试题id矩阵
            per: 考点分布数组
            """
            print('\ncurrent:', e_type)
            current_rate = 0
            flag = False
            already_me = already[e_type]
            print('已有试题组合：', already_me)

            # max_row为章节总数,max_col为拥有最多试题的章节的试题数
            max_row = len(secs)
            max_col = max(secs_length)
            id_matrix = np.zeros((max_row, max_col), dtype=int)
            for row in range(max_row):
                id_matrix[row, :len(id_list[row])] = id_list[row]
            print('试题id矩阵:\n', id_matrix)

            # 计算试题数矩阵
            weight = np.array(per, dtype=float)
            num = np.floor(weight * count)
            print('当前分布下的各章节试题数矩阵:\n', num)
            margin = count - np.sum(num, dtype=int)
            num[np.argmax(num)] += margin
            print('误差补正后的各章节试题数矩阵:\n', num)
            mask = np.zeros([max_row, max_col], dtype=int)
            try:
                for iter in range(100):
                    for i in range(len(secs_length)):
                        # sec_l为该章节的试题总数
                        sec_l = secs_length[i]
                        # x为该章节需要抽取的试题数
                        x = int(num[i])
                        rand = np.zeros(sec_l, dtype=int)
                        one = np.ones(x, dtype=int)
                        rand[:one.shape[0]] = one
                        np.random.shuffle(rand)
                        mask[i, :rand.shape[0]] = rand
                    print('蒙版矩阵:\n', mask)
                    mul = np.extract(mask, id_matrix)
                    print('蒙版结果:\n', mul)
                    current_list = mul.tolist()
                    print('列表输出：\n', current_list)
                    if already_me != []:
                        print('当前试卷为课程的后续试卷')
                        for already_list in already_me:  # 随机组合在已有试题组内进行重复率比较
                            print('current_list:', current_list)
                            print('already_list:', already_list)
                            current_rate = len(set(current_list) &
                                               set(already_list)) / count
                            if current_rate <= 0.25:
                                flag = True
                                continue
                            else:
                                print('随机失败, current_rate:', current_rate)
                                flag = False
                                break
                        if flag:  # 通过全部已有试题组合的重复率匹配后退出循环
                            break
                        else:
                            continue
                    else:
                        print('当前试卷为课程的首张试卷')
                        current_rate = 0
                        break
                print('随机完毕, current_rate:', current_rate)
                return current_list, current_rate

            except ValueError as e:
                print(e)
                print('某些章节设定比重过高，对应试题量不足，请扩充题库')
                return [], 0

    def case1(courseID, e_type, count, per):
        secs_length = []
        id_list = []
        secs = []
        per = []
        for sec in percentage:
            secs.append(sec)
            per.append(percentage[sec])
        for sec in secs:
            choices = Choice.objects.filter(courseID=courseID,
                                            is_single=True, sec=sec).values_list('id', flat=True)
            id_list.append(list(choices))
            length = len(choices)
            secs_length.append(length)
        print('所有章节的试题数列表:\n', secs_length)
        print('所有章节的试题列表:\n', id_list)

        es[e_type], rpr[e_type] = pickle_in_random(
            secs_length=secs_length, secs=secs, count=count, e_type=e_type, id_list=id_list, per=per)
        return es[e_type], rpr[e_type]

    def case2(courseID, e_type, count, per):
        secs_length = []
        id_list = []
        secs = []
        per = []
        for sec in percentage:
            secs.append(sec)
            per.append(percentage[sec])
        for sec in secs:
            choices = Choice.objects.filter(courseID=courseID,
                                            is_single=False, sec=sec).values_list('id', flat=True)
            id_list.append(list(choices))
            length = len(choices)
            secs_length.append(length)
        print('所有章节的试题数列表:\n', secs_length)
        print('所有章节的试题列表:\n', id_list)

        es[e_type], rpr[e_type] = pickle_in_random(
            secs_length=secs_length, secs=secs, count=count, e_type=e_type, id_list=id_list, per=per)
        return es[e_type], rpr[e_type]

    def case3(courseID, e_type, count, per):
        secs_length = []
        id_list = []
        secs = []
        per = []
        for sec in percentage:
            secs.append(sec)
            per.append(percentage[sec])
        for sec in secs:
            judges = Judge.objects.filter(
                courseID=courseID, sec=sec).values_list('id', flat=True)
            id_list.append(list(judges))
            length = len(judges)
            secs_length.append(length)
        print('所有章节的试题数列表:\n', secs_length)
        print('所有章节的试题列表:\n', id_list)

        es[e_type], rpr[e_type] = pickle_in_random(
            secs_length=secs_length, secs=secs, count=count, e_type=e_type, id_list=id_list, per=per)
        return es[e_type], rpr[e_type]

    def case4(courseID, e_type, count, per):
        secs_length = []
        id_list = []
        secs = []
        per = []
        for sec in percentage:
            secs.append(sec)
            per.append(percentage[sec])
        for sec in secs:
            s_answers = S_answer.objects.filter(
                courseID=courseID, sec=sec).values_list('id', flat=True)
            id_list.append(list(s_answers))
            length = len(s_answers)
            secs_length.append(length)
        print('所有章节的试题数列表:\n', secs_length)
        print('所有章节的试题列表:\n', id_list)

        es[e_type], rpr[e_type] = pickle_in_random(
            secs_length=secs_length, secs=secs, count=count, e_type=e_type, id_list=id_list, per=per)
        return es[e_type], rpr[e_type]

    def case5(courseID, e_type, count, per):
        secs_length = []
        id_list = []
        secs = []
        per = []
        for sec in percentage:
            secs.append(sec)
            per.append(percentage[sec])
        for sec in secs:
            blanks = Blank.objects.filter(
                courseID=courseID, sec=sec).values_list('id', flat=True)
            id_list.append(list(blanks))
            length = len(blanks)
            secs_length.append(length)
        print('所有章节的试题数列表:\n', secs_length)
        print('所有章节的试题列表:\n', id_list)

        es[e_type], rpr[e_type] = pickle_in_random(
            secs_length=secs_length, secs=secs, count=count, e_type=e_type, id_list=id_list, per=per)
        return es[e_type], rpr[e_type]

    switch = {
        's-choice': case1,
        'm-choice': case2,
        'judge': case3,
        's-answer': case4,
        'blank': case5,
    }
    # 取回题目描述

    def extra_case1(id_list):
        id_desc = {}
        querysets = Choice.objects.filter(
            id__in=id_list).values('id', 'descri')
        for one in querysets:
            id_desc[one['id']] = one['descri']
        return id_desc

    def extra_case2(id_list):
        id_desc = {}
        querysets = Choice.objects.filter(
            id__in=id_list).values('id', 'descri')
        for one in querysets:
            id_desc[one['id']] = one['descri']
        return id_desc

    def extra_case3(id_list):
        id_desc = {}
        querysets = Judge.objects.filter(
            id__in=id_list).values('id', 'descri')
        for one in querysets:
            id_desc[one['id']] = one['descri']
        return id_desc

    def extra_case4(id_list):
        id_desc = {}
        querysets = S_answer.objects.filter(
            id__in=id_list).values('id', 'descri')
        for one in querysets:
            id_desc[one['id']] = one['descri']
        return id_desc

    def extra_case5(id_list):
        id_desc = {}
        querysets = Blank.objects.filter(
            id__in=id_list).values('id', 'descri')
        for one in querysets:
            id_desc[one['id']] = one['descri']
        return id_desc

    switch_extra = {
        's-choice': extra_case1,
        'm-choice': extra_case2,
        'judge': extra_case3,
        's-answer': extra_case4,
        'blank': extra_case5,
    }

    data = request.body.decode('utf-8')
    json_request = json.loads(data)
    courseID = json_request['courseID']
    e_dict = json_request['dict']
    e_percentage = json_request['percentage']

    drafts = Draft.objects.filter(
        courseID=courseID).values_list('draft_string', flat=True)
    already = {'s-choice': [], 'm-choice': [],
               'judge': [], 's-answer': [], 'blank': []}

    start = time.time()
    for d in drafts:
        raw = json.loads(d)
        print(raw)
        for key in raw:
            already[key].append(map(eval, list(raw[key])))
            print(list(already[key]))
    print('已有试题组合：', already)

    es = {}
    rpr = {}
    for e_type in e_dict:
        count = int(e_dict[e_type])
        percentage = e_percentage[e_type]
        es[e_type], rpr[e_type] = switch[e_type](
            courseID, e_type, count, percentage)
    print('随机结果', es)
    end = time.time()
    print('随机执行时间：', end - start)
    # 取回题目描述信息
    details = {}
    for e_type in es:
        if es[e_type] == 0:
            details[e_type] = False
        else:
            details[e_type] = switch_extra[e_type](es[e_type])
    result = {'result': details, 'rpr': rpr}
    return JsonResponse(result)


class exam_build(Auth, LoginRequiredMixin, View):
    """
    POST方法创建试卷对象
    """

    def get(self, request):
        courses = Course.objects.filter(teacher=request.user.user_id)
        return render(request, 'exam_build.html', {'courses': courses})

    def post(self, request):

        def case1(distrib):
            for e_id in distrib:
                point = distrib[e_id]
                obj = Choice.objects.get(id=e_id)
                obj.point = point
                obj.save()
                exam.choices.add(obj)

        def case2(distrib):
            for e_id in distrib:
                point = distrib[e_id]
                obj = Choice.objects.get(id=e_id)
                obj.point = point
                obj.save()
                exam.choices.add(obj)

        def case3(distrib):
            for e_id in distrib:
                point = distrib[e_id]
                obj = Judge.objects.get(id=e_id)
                obj.point = point
                obj.save()
                exam.judges.add(obj)

        def case4(distrib):
            for e_id in distrib:
                point = distrib[e_id]
                obj = S_answer.objects.get(id=e_id)
                obj.point = point
                obj.save()
                exam.s_answers.add(obj)

        def case5(distrib):
            for e_id in distrib:
                point = distrib[e_id]
                obj = Blank.objects.get(id=e_id)
                obj.point = point
                obj.save()
                exam.blanks.add(obj)

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }

        print(request.POST)
        data = request.body.decode('utf-8')
        json_request = json.loads(data)
        print(json_request)
        e_dict = json_request['dict']
        form = ExamForm(data=json_request)
        if form.is_valid():
            exam = form.save()
            for e_type in e_dict:
                switch[e_type](e_dict[e_type])
            print('完成试卷创建！试卷id：', exam.id)
            result = {'result': True}

            form = DraftForm(data=json_request)
            if form.is_valid():
                draft = form.save(commit=False)
                e_string = json.dumps(e_dict)
                draft.draft_string = e_string
                draft.exam = exam
                draft.save()
                print('完成快照创建！快照id：', draft.id)
            else:
                print(form.errors)

        else:
            print(form.errors)
            result = {'result': False}
        return JsonResponse(result)


def exam_delete(request, courseID):
    if request.method == 'POST':
        exam_id = request.POST['exam_id']
        exam = Exam.objects.get(id=exam_id)
        exam.delete()
        result = {'result': True}
        return JsonResponse(result)


class exam_manage(Auth, LoginRequiredMixin, View):
    """
    POST方法创建docx试卷
    """

    def get(self, request, courseID):
        exams = Exam.objects.prefetch_related(
            'courseID').filter(courseID=courseID)
        return render(request, 'exam_manage.html', {'exams': exams})

    def post(self, request, courseID):

        def convert_to_pdf(doc):
            """
            docx转pdf
            """
            if client==None:
                return doc2pdf_linux(doc)
            try:
                pythoncom.CoInitialize()
                word = client.CreateObject("Word.Application")
                new_name = doc.replace(".docx", r".pdf")
                worddoc = word.Documents.Open(doc)
                worddoc.SaveAs(new_name, FileFormat=17)
                worddoc.Close()
                pdf_path = path.replace(".docx", r".pdf")
            except Exception as e:
                print(e)
                pdf_path = False
            finally:
                word.Quit()
                return pdf_path

        def doc2pdf_linux(doc):
            try:
                cmd =[ "libreoffice","--convert-to pdf", doc]
                p = subprocess.Popen(cmd, stderr=subprocess.PIPE, stdout=subprocess.PIPE)
                p.wait(timeout=10)
                return path.replace('.docx',r'.pdf')
            except Exception as e:
                raise e

        def cal_points(querysets):
            try:
                points = list(querysets.values_list('point'))
                points_arr = np.array(points)
                points_addup = np.sum(points_arr)
                return points_addup
            except IndexError:
                pass
        exam_id = request.POST['exam_id']
        exam_format = request.POST['format']
        exam = Exam.objects.get(id=exam_id)
        print(request.POST)
        print(exam.choices.all())
        s_choices = exam.choices.filter(is_single=True).order_by('?')
        m_choices = exam.choices.filter(is_single=False).order_by('?')
        blanks = exam.blanks.all().order_by('?')
        judges = exam.judges.all().order_by('?')
        s_answers = exam.s_answers.all().order_by('?')

        e_points = {
            's_choices': cal_points(s_choices),
            'm_choices': cal_points(m_choices),
            'blanks': cal_points(blanks),
            'judges': cal_points(judges),
            's_answers': cal_points(s_answers),
        }
        print(e_points)
        context = {
            'exam': exam,
            'e_points': e_points,
            's_choices': s_choices,
            'm_choices': m_choices,
            'blanks': blanks,
            'judges': judges,
            's_answers': s_answers,
        }

        doc = DocxTemplate('media/exam_template.docx')
        doc.render(context)
        path = 'media/user_%s/%s.docx' % (request.user.user_id, exam.title)
        doc.save(path)

        if exam_format == 'docx':
            result = {'title': exam.title, 'path': path}
        elif exam_format == 'pdf':
            new_doc = os.path.abspath(path)
            pdf_path = convert_to_pdf(new_doc)
            result = {'title': exam.title, 'path': pdf_path}
        else:
            result = {'title': '预览试卷', 'context_json': context_json}
        return JsonResponse(result)


@is_teacher
def exam_retrieve(request):
    if request.method == 'GET':

        def case1():
            choices = Choice.objects.filter(id=e_id).values(
                'descri', 'A', 'B', 'C', 'D', 'E', 'F', 'answer', 'sec', 'courseID')
            return choices

        def case2():
            choices = Choice.objects.filter(id=e_id).values(
                'descri', 'A', 'B', 'C', 'D', 'E', 'F', 'answer', 'sec', 'courseID')
            return choices

        def case3():
            judges = Judge.objects.filter(id=e_id).values(
                'descri', 'answer', 'sec', 'courseID')
            return judges

        def case4():
            s_answers = S_answer.objects.filter(
                id=e_id).values('descri', 'sec', 'answer')
            return s_answers

        def case5():
            blanks = Blank.objects.filter(id=e_id).values(
                'descri', 'blank1', 'blank2', 'blank3', 'blank4', 'blank5', 'sec', 'courseID')
            return blanks

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }

        e_type = request.GET['e_type']
        e_id = request.GET['id']
        querysetsV = switch[e_type]()
        item = list(querysetsV)
        return JsonResponse(item, safe=False)


class bank_manage(Auth, LoginRequiredMixin, View):
    """
    GET方法返回课程列表，POST方法用于批量删除试题
    """

    def get(self, request):
        courses = Course.objects.filter(
            teacher=request.user.user_id).values('courseID', 'course')
        return render(request, 'bank_manage.html', {'courses': courses})

    def post(self, request):
        if request.method == 'POST':

            def case1(e_list):
                for e_id in e_list:
                    print(e_id)
                    Choice.objects.filter(id=e_id).delete()
                return True

            def case2(e_list):
                for e_id in e_list:
                    Choice.objects.filter(id=e_id).delete()
                return True

            def case3(e_list):
                for e_id in e_list:
                    Judge.objects.filter(id=e_id).delete()
                return True

            def case4(e_list):
                for e_id in e_list:
                    S_answer.objects.filter(id=e_id).delete()
                return True

            def case5(e_list):
                for e_id in e_list:
                    Blank.objects.filter(id=e_id).delete()
                return True

            switch = {
                's-choice': case1,
                'm-choice': case2,
                'judge': case3,
                's-answer': case4,
                'blank': case5,
            }
            data = request.body.decode('utf-8')
            json_request = json.loads(data)
            e_type = json_request['e_type']
            try:
                e_list = json_request['e_list']
                result = {'result': switch[e_type](e_list)}
            except MultiValueDictKeyError:
                result = {'result': False}
            return JsonResponse(result)


class bank_bulk(Auth, LoginRequiredMixin, View):

    def post(self, request):

        def case1():
            WorkLish = []
            for n in range(1, rows_num):
                row = sheet.row_values(n)
                # 批量写入数据库， 需定制
                WorkLish.append(Choice(courseID=course, is_single=True, descri=row[0], A=row[
                                1], B=row[2], C=row[3], D=row[4], E=row[5], F=row[6], answer=row[7], sec=row[8]))
            Choice.objects.bulk_create(WorkLish)
            return {"result": True}

        def case2():
            WorkLish = []
            for n in range(1, rows_num):
                row = sheet.row_values(n)
                # 批量写入数据库， 需定制
                WorkLish.append(Choice(courseID=course, is_single=False, descri=row[0], A=row[
                                1], B=row[2], C=row[3], D=row[4], E=row[5], F=row[6], answer=row[7], sec=row[8]))
            Choice.objects.bulk_create(WorkLish)
            return {"result": True}

        def case3():
            WorkLish = []
            for n in range(1, rows_num):
                row = sheet.row_values(n)
                # 批量写入数据库， 需定制
                WorkLish.append(
                    Judge(courseID=course, descri=row[0], answer=row[1], sec=row[2]))
            Judge.objects.bulk_create(WorkLish)
            return {"result": True}

        def case4():
            WorkLish = []
            for n in range(1, rows_num):
                row = sheet.row_values(n)
                # 批量写入数据库， 需定制
                WorkLish.append(
                    S_answer(courseID=course, descri=row[0], answer=row[1], sec=row[2]))
            S_answer.objects.bulk_create(WorkLish)
            return {"result": True}

        def case5():
            WorkLish = []
            for n in range(1, rows_num):
                row = sheet.row_values(n)
                # 批量写入数据库， 需定制
                WorkLish.append(Blank(courseID=course, descri=row[0], blank1=row[1], blank2=row[
                                2], blank3=row[3], blank4=row[4], blank5=row[5], blank6=row[6], sec=row[7]))
            Blank.objects.bulk_create(WorkLish)
            return {"result": True}

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }

        form = BankProcessForm(data=request.POST, files=request.FILES)

        if form.is_valid():
            try:
                raw_bank = form.save()
                course = Course.objects.get(courseID=raw_bank.courseID)
                # 获取硬盘中的绝对地址
                path = raw_bank.file.path
                print(path)
                # 操作完不需要手动关闭
                excel = xlrd.open_workbook(
                    filename=str(path), encoding_override='utf8')
                # 使用第一张表
                sheet = excel.sheet_by_index(0)
                rows_num = sheet.nrows
                result = switch[form.cleaned_data['e_type']]()
            except xlrd.biffh.XLRDError as e:
                result = {'result': 'Do not support'}
        else:
            result = {'result': 'not valid form'}
            print(form)
        return JsonResponse(result)


@is_teacher
def bank_query(request):

    if request.method == 'POST':

        def case1(courseID, keyword):
            choices = Choice.objects.filter(
                courseID=courseID, descri__icontains=keyword, is_single=True).values('id', 'descri')
            return choices

        def case2(courseID, keyword):
            choices = Choice.objects.filter(
                courseID=courseID, descri__icontains=keyword, is_single=False).values('id', 'descri')
            return choices

        def case3(courseID, keyword):
            judges = Judge.objects.filter(
                courseID=courseID, descri__icontains=keyword).values('id', 'descri')
            return judges

        def case4(courseID, keyword):
            s_answers = S_answer.objects.filter(
                courseID=courseID, descri__icontains=keyword).values('id', 'descri')
            return s_answers

        def case5(courseID, keyword):
            blanks = Blank.objects.filter(
                courseID=courseID, descri__icontains=keyword).values('id', 'descri')
            return blanks

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }
        print(request.POST)
        courseID = request.POST['courseID']
        e_type = request.POST['e_type']
        keyword = request.POST['keyword']
        items = switch[e_type](courseID, keyword)
        result = list(items)
        return JsonResponse(result, safe=False)


@is_teacher
def bank_query_alter(request, template_name='partials/e_candidate.html'):

    if request.method == 'POST':

        def case1(courseID, keyword):
            choices = Choice.objects.filter(
                courseID=courseID, descri__icontains=keyword, is_single=True)
            return choices

        def case2(courseID, keyword):
            choices = Choice.objects.filter(
                courseID=courseID, descri__icontains=keyword, is_single=False)
            return choices

        def case3(courseID, keyword):
            judges = Judge.objects.filter(
                courseID=courseID, descri__icontains=keyword)
            return judges

        def case4(courseID, keyword):
            s_answers = S_answer.objects.filter(
                courseID=courseID, descri__icontains=keyword)
            return s_answers

        def case5(courseID, keyword):
            blanks = Blank.objects.filter(
                courseID=courseID, descri__icontains=keyword)
            return blanks

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }

        courseID = request.POST['courseID']
        e_type = request.POST['e_type']
        keyword = request.POST['keyword']
        items = switch[e_type](courseID, keyword)
        return render(request, template_name, {'items': items, 'e_type': e_type})


@is_teacher
def bank_add(request, e_type):
    """
    添加单条试题
    """
    if request.method == 'GET':
        def case1():
            form = ChoiceForm()
            return form

        def case2():
            form = ChoiceForm()
            return form

        def case3():
            form = JudgeForm()
            return form

        def case4():
            form = S_answerForm()
            return form

        def case5():
            form = BlankForm()
            return form

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }
        form = switch[e_type]()

        return HttpResponse(str(form))

    if request.method == 'POST':
        def case1():
            form = ChoiceForm(data=request.POST)
            return form

        def case2():
            form = ChoiceForm(data=request.POST)
            return form

        def case3():
            form = JudgeForm(data=request.POST)
            return form

        def case4():
            form = S_answerForm(data=request.POST)
            return form

        def case5():
            form = BlankForm(data=request.POST)
            return form

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }
        print(request.POST)
        form = switch[e_type]()
        if form.is_valid():
            form.save()
            result = {'result': True}
        else:
            result = {'result': False}
            print(form.errors)
        return JsonResponse(result)


@is_teacher
def bank_edit(request):
    """
    编辑单条试题
    """
    if request.method == 'GET':
        def case1():
            form = ChoiceForm()
            return form

        def case2():
            form = ChoiceForm()
            return form

        def case3():
            form = JudgeForm()
            return form

        def case4():
            form = S_answerForm()
            return form

        def case5():
            form = BlankForm()
            return form

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }
        e_type = request.GET['e_type']
        form = switch[e_type]()

        return HttpResponse(str(form))

    if request.method == 'POST':

        def case1(e_id):
            form = ChoiceForm(data=request.POST)
            if form.is_valid():
                choice = Choice.objects.get(id=e_id)
                choice.descri = form.cleaned_data['descri']
                choice.A = form.cleaned_data['A']
                choice.B = form.cleaned_data['B']
                choice.C = form.cleaned_data['C']
                choice.D = form.cleaned_data['D']
                choice.E = form.cleaned_data['E']
                choice.F = form.cleaned_data['F']
                choice.answer = form.cleaned_data['answer']
                choice.sec = form.cleaned_data['sec']
                choice.save()
                result = {'result': True}
            else:
                result = {'result': False}
                print(form.errors)
            return result

        def case2(e_id):
            form = ChoiceForm(data=request.POST)
            if form.is_valid():
                choice = Choice.objects.get(id=e_id)
                choice.descri = form.cleaned_data['descri']
                choice.A = form.cleaned_data['A']
                choice.B = form.cleaned_data['B']
                choice.C = form.cleaned_data['C']
                choice.D = form.cleaned_data['D']
                choice.E = form.cleaned_data['D']
                choice.F = form.cleaned_data['F']
                choice.answer = form.cleaned_data['answer']
                choice.sec = form.cleaned_data['sec']
                choice.save()
                print(form)
                result = {'result': True}
            else:
                result = {'result': False}
                print(form.errors)
            return result

        def case3(e_id):
            form = JudgeForm(data=request.POST)
            if form.is_valid():
                judge = Judge.objects.get(id=e_id)
                judge.descri = form.cleaned_data['descri']
                judge.answer = form.cleaned_data['answer']
                judge.sec = form.cleaned_data['sec']
                print(form.cleaned_data['answer'])
                print(form)
                judge.save()
                result = {'result': True}
            else:
                result = {'result': False}
                print(form.errors)
            return result

        def case4(e_id):
            form = S_answerForm(data=request.POST)
            if form.is_valid():
                s_answer = S_answer.objects.get(id=e_id)
                s_answer.descri = form.cleaned_data['descri']
                s_answer.answer = form.cleaned_data['answer']
                s_answer.sec = form.cleaned_data['sec']
                s_answer.save()
                result = {'result': True}
            else:
                result = {'result': False}
                print(form.errors)
            return result

        def case5(e_id):
            form = BlankForm(data=request.POST)
            if form.is_valid():
                blank = Blank.objects.get(id=e_id)
                blank.descri = form.cleaned_data['descri']
                blank.blank1 = form.cleaned_data['blank1']
                blank.blank2 = form.cleaned_data['blank2']
                blank.blank3 = form.cleaned_data['blank3']
                blank.blank4 = form.cleaned_data['blank4']
                blank.blank5 = form.cleaned_data['blank5']
                blank.blank6 = form.cleaned_data['blank6']
                blank.sec = form.cleaned_data['sec']
                blank.save()
                result = {'result': True}
            else:
                result = {'result': False}
                print(form.errors)
            return result

        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }
        print(request.POST)
        e_type = request.POST['e_type']
        e_id = request.POST['id']
        result = switch[e_type](e_id)
        return JsonResponse(result)


@is_teacher
def bank_drop(request):
    if request.method == 'POST':
        data = request.body.decode('utf-8')
        json_request = json.loads(data)
        courseID = json_request['courseID']
        try:
            Choice.objects.filter(courseID=courseID).delete()
            Judge.objects.filter(courseID=courseID).delete()
            S_answer.objects.filter(courseID=courseID).delete()
            Blank.objects.filter(courseID=courseID).delete()
            result = {'result': True}
        except:
            result = {'result': False}
        return JsonResponse(result)


@is_teacher
def bank_export(request):

    def case1():
        choices = Choice.objects.filter(
            courseID=courseID, is_single=True).values()
        return choices

    def case2():
        choices = Choice.objects.filter(
            courseID=courseID, is_single=False).values()
        return choices

    def case3():
        judges = Judge.objects.filter(courseID=courseID).values()
        return judges

    def case4():
        s_answers = S_answer.objects.filter(courseID=courseID).values()
        return s_answers

    def case5():
        blanks = Blank.objects.filter(courseID=courseID).values()
        return blanks

    def set_style(name, height, bold=False):
        style = xlwt.XFStyle()  # 初始化样式

        font = xlwt.Font()  # 为样式创建字体
        font.name = name
        font.bold = bold
        font.color_index = 4
        font.height = height

        style.font = font
        return style

    def get_row0(queryset):
        raw_data = queryset
        raw_data.pop('created_at')
        raw_data.pop('courseID_id')
        raw_data.pop('id')
        raw_data.pop('point')
        try:
            raw_data.pop('is_single')
        finally:
            row0 = []
            for key in raw_data:
                row0.append(key)
            return row0

    def clean_raw_data(queryset):
        raw_data = queryset
        raw_data.pop('created_at')
        raw_data.pop('courseID_id')
        raw_data.pop('id')
        raw_data.pop('point')
        try:
            raw_data.pop('is_single')
        finally:
            return raw_data

    def write_excel(e_type, querysets):
        # 创建工作簿
        workbook = xlwt.Workbook(encoding='utf-8')
        # 创建sheet
        data_sheet = workbook.add_sheet('Sheet1')
        row0 = get_row0(querysets[0])
        # 生成第一行
        for i in range(len(row0)):
            data_sheet.write(0, i, row0[i], set_style(
                'Microsoft Yahei', 220, False))

        for n in range(querysets.count()):
            data = clean_raw_data(querysets[n])
            for i in range(len(row0)):
                data_sheet.write(n + 1, i, data[row0[i]], set_style(
                    'Microsoft Yahei', 220, False))
        # 保存文件
        workbook.save('media/user_%s/%s.xls' % (request.user.user_id, e_type))

    def export(e_type, querysets):
        try:
            write_excel(e_type, querysets)
            return True
        except IndexError:
            return False

    if request.method == 'POST':
        result = {}
        courseID = request.POST['courseID']
        switch = {
            's-choice': case1,
            'm-choice': case2,
            'judge': case3,
            's-answer': case4,
            'blank': case5,
        }

        for key in switch:
            querysets = switch[key]()
            result[key] = export(key, querysets)

        return JsonResponse(result)


class extract_from_docx(Auth, LoginRequiredMixin, View):

    def get(self, request):
        return render(request, 'docx_2_xls.html')

    def post(self, request):

        def set_style(name, height, bold=False):
            style = xlwt.XFStyle()  # 初始化样式

            font = xlwt.Font()  # 为样式创建字体
            font.name = name
            font.bold = bold
            font.color_index = 4
            font.height = height

            style.font = font
            return style

        def docx_2_xls(e_type_name, pre_file):
            try:
                # 读取模板
                template = 'media/%s_template.xls' % e_type_name
                workbook_temp = xlrd.open_workbook(template)
                sheet_temp = workbook_temp.sheet_by_name('Sheet1')
                row0 = sheet_temp.row_values(0)

                # 创建工作簿
                workbook = xlwt.Workbook(encoding='utf-8')
                # 创建sheet
                data_sheet = workbook.add_sheet('Sheet1')
                # 生成第一行
                for i in range(len(row0)):
                    data_sheet.write(0, i, row0[i], set_style(
                        'Microsoft Yahei', 220, False))

                # 读取docx预处理文件
                doc = docx.Document(pre_file)
                text = ''
                for p in doc.paragraphs:
                    p.text = re.sub((r'[0-9A-Z]+[、.．] *'), ' ',
                                    p.text)  # 替换选项标号和题目标号
                    p.text = re.sub((r' *（ *） *'), '（）', p.text)  # 替换中文带空格括号
                    text = text + ' ' + p.text
                #print(text)
                es = re.finditer((r"\{(.*?)\}"), text)  # 匹配预处理目标

                # 写入到新的xls模板中
                row = 1
                for match in es:
                    # split()默认分隔符为所有的空字符，包括空格、换行(\n)、制表符(\t)等
                    e = match.group(1).split()
                    col = 0
                    for i in e:
                        if i:
                            data_sheet.write(row, col, i, set_style(
                                'Microsoft Yahei', 220, False))
                            col = col + 1
                        else:
                            pass
                    row = row + 1

                # 保存文件
                path = 'media/user_%s/%s_格式化题库.xls' % (
                    request.user.user_id, e_type_name)
                workbook.save(path)
                errors = 'clear'
                return path, errors
            except FileNotFoundError:
                path = 'media/%s_template.xls' % e_type_name
                errors = 'wrong filename'
                return path, errors

        form = DocxForm(request.POST, request.FILES)
        if form.is_valid():
            pre_file = request.FILES.get('file')
            e_type_name = pre_file.name.split('.')[0]
            print(e_type_name)
            path, errors = docx_2_xls(e_type_name, pre_file)
            result = {'e_type_name': e_type_name,
                      'path': path, 'errors': errors}
            return JsonResponse(result)
